import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'D:\Source mobile 2\BTLMobile2\TàiLiệu\baocao8.docx')

print("Mapping images to paragraphs and captions:")
image_map = []
for p_idx, p in enumerate(doc.paragraphs):
    # Check for drawing/blip
    images_in_p = []
    for r in p.runs:
        element = r._r
        blips = element.xpath('.//a:blip')
        for b in blips:
            rId = b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if rId:
                target = doc.part.rels[rId].target_ref
                images_in_p.append((rId, target))
    if images_in_p:
        # Check adjacent text
        caption = ""
        if p_idx + 1 < len(doc.paragraphs):
            caption = doc.paragraphs[p_idx + 1].text.strip()
        image_map.append({
            'p_idx': p_idx,
            'images': images_in_p,
            'caption': caption,
            'p_text': p.text.strip()
        })

print(f"Total image paragraphs found: {len(image_map)}")
for item in image_map:
    imgs = ", ".join([f"{t} ({rId})" for rId, t in item['images']])
    print(f"\n[P{item['p_idx']}] Img: {imgs}")
    print(f"Caption: {item['caption']}")
