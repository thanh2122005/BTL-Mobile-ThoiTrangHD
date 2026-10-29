import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'D:\Source mobile 2\BTLMobile2\TàiLiệu\baocao8.docx'
doc = docx.Document(doc_path)

print(f"=== BÁO CÁO: {os.path.basename(doc_path)} ===")
print(f"Tổng số đoạn văn (paragraphs): {len(doc.paragraphs)}")
print(f"Tổng số bảng (tables): {len(doc.tables)}")

headings = []
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if p.style.name.startswith('Heading') or (text.isupper() and len(text) < 100 and len(text) > 3):
        headings.append((idx, p.style.name, text))

print("\n--- CÁC TIÊU ĐỀ CHÍNH ---")
for idx, style, text in headings[:50]:
    print(f"[{idx}][{style}] {text}")

# Check trường, thầy, khoa
print("\n--- KIỂM TRA THÔNG TIN MỞ ĐẦU ---")
for i in range(min(40, len(doc.paragraphs))):
    t = doc.paragraphs[i].text.strip()
    if t:
        print(f"P{i}: {t}")

# Check images and captions
print("\n--- DANH SÁCH HÌNH ẢNH & CHÚ THÍCH HÌNH ---")
figure_captions = []
for idx, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.lower().startswith('hình ') or t.lower().startswith('hình:') or 'giao diện' in t.lower() or 'sơ đồ' in t.lower():
        figure_captions.append((idx, t))

print(f"Tìm thấy {len(figure_captions)} chú thích hình:")
for idx, cap in figure_captions:
    print(f"  Line {idx}: {cap}")
